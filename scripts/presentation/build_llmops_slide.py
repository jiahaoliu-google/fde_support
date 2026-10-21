"""Builds docs/llmops_slide.pptx: one editable 16:9 slide on LLMOps, the evaluation framework
and observability. Reuses the drawing helpers and design tokens from build_cost_slide.py.

    python scripts/presentation/build_llmops_slide.py # requires python-pptx==1.0.2

Content source: docs/EVALUATION_PLAN.md (KPIs, golden set, gates) and backend/ (OpenTelemetry,
Model Armor, circuit breaker).
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from build_cost_slide import C, MONO, SANS, background, fill_frame, px, rgb, set_cell_borders, shape, text

OUT = Path(__file__).resolve().parent.parent.parent / "docs" / "llmops_slide.pptx"

GOLDEN_SET = [  # (bucket, cases, colour, what it tests)
    ("Search-ready, in corpus", 45, "8B5CF6", "grounded answers"),
    ("Urgent / frustrated", 25, "FB7185", "human handoff"),
    ("Vague, single turn", 15, "FBBF24", "clarify, no search"),
    ("Multi-turn, vague \u2192 clear", 10, "22D3EE", "gating over turns"),
    ("Out of corpus", 5, "94A3B8", "no hallucination"),
]
SCORERS = [
    ("Deterministic", "F1, ROUGE-L, BLEU, search gating, latency, tokens"),
    ("LLM judge", "Gemini Flash: grounding recall, faithfulness"),
    ("Dynamic rubrics", "Per-case criteria for clarifying questions"),
]
GATES = [  # (KPI, target, gate)
    ("Routing accuracy", "\u2265 95%", "PR"),
    ("Urgent / frustrated F1", "\u2265 0.95", "PR"),
    ("Urgent tickets bypass RAG", "100%", "PR"),
    ("Search gating (vague / specific)", "\u2265 95%", "PR"),
    ("ROUGE-L vs. brand reply", "\u2265 0.35", "PR"),
    ("BLEU-4 vs. brand reply", "\u2265 0.30", "PR"),
    ("API error / timeout rate", "0%", "PR"),
    ("Grounding recall (key facts)", "\u2265 85%", "Nightly"),
    ("RAG faithfulness", "\u2265 90%", "Nightly"),
    ("P95 end-to-end latency", "\u2264 3.0 s", "PR"),
]
OBSERVABILITY = [  # (title, detail, colour)
    ("Distributed tracing", "OpenTelemetry (FastAPI + httpx) \u2192 Cloud Trace: P95 / P99 per hop", "38BDF8"),
    ("Structured logs \u2192 BigQuery", "Cloud Logging sink feeds Agent Analytics: routes, tokens, cost", "22D3EE"),
    ("Online quality sampling", "Async LLM judge on sampled live traces + escalation rate, CSAT", "C084FC"),
    ("PII-safe telemetry", "Model Armor / Cloud DLP redacts input before agents and logs", "34D399"),
    ("Resilience signals", "Circuit breaker (3 failures / 30 s) and 4xx, 5xx, 429 counters", "FBBF24"),
    ("Alerting", "SLO breaches page on-call (PagerDuty)", "FB7185"),
]
LIFECYCLE = [  # (step, detail)
    ("Change", "Prompt, model, tool, code"),
    ("PR gate", "Deterministic KPIs"),
    ("Nightly", "LLM judges + rubrics"),
    ("Deploy", "Cloud Build \u2192 Cloud Run"),
    ("Monitor", "Traces, logs, alerts"),
    ("Learn", "Failures \u2192 golden set"),
]


def header(slide):
    text(slide, 72, 52, 1100, 26, [[("FDE SUPPORT PLATFORM \u00b7 OPERATIONS", 17, C["pro"], True, SANS, 300)]])
    text(slide, 72, 82, 1300, 74, [[
        ("LLMOps, ", 54, C["white"], True),
        ("Evaluation Framework & Observability", 54, C["lav"], True),
    ]], wrap=False)
    text(slide, 72, 150, 1300, 28, [[("Every release is gated on a 100-ticket golden set; production is watched by traces, logs and sampled judges.", 19, C["muted"])]])
    b = shape(slide, 1448, 112, 400, 40, C["warn_bg"], "4A3F1A", radius=20)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_frame(tf, [[("Status: framework built \u00b7 baseline run pending", 15, C["warn"], True)]], PP_ALIGN.CENTER)


def panel(slide, x, y, w, h, num, title, sub):
    shape(slide, x, y, w, h, C["panel"], C["border"], radius=16)
    n = shape(slide, x + 28, y + 24, 34, 34, "1C1830", None, radius=9)
    tf = n.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_frame(tf, [[(num, 17, C["pro"], True)]], PP_ALIGN.CENTER)
    text(slide, x + 76, y + 24, w - 100, 34, [[(title, 24, C["text"], True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 28, y + 68, w - 56, 24, [[(sub, 16, C["muted"])]])


def eval_panel(slide, x, y, w, h):
    panel(slide, x, y, w, h, "1", "Offline evaluation", "Golden set: 100 tickets sourced from the RAG corpus")
    ix, iw = x + 28, w - 56
    bx, by = ix, y + 108
    total = sum(n for _, n, _, _ in GOLDEN_SET)
    for i, (_, n, color, _) in enumerate(GOLDEN_SET):
        seg = iw * n / total if i < len(GOLDEN_SET) - 1 else ix + iw - bx
        shape(slide, bx, by, seg, 36, color)
        bx += seg
    ry = by + 52
    for label, n, color, purpose in GOLDEN_SET:
        shape(slide, ix, ry + 9, 14, 14, color, radius=3)
        text(slide, ix + 26, ry, 250, 32, [[(label, 17.5, C["text"])]], anchor=MSO_ANCHOR.MIDDLE)
        text(slide, ix + 270, ry, iw - 320, 32, [[(purpose, 15, C["muted"])]], anchor=MSO_ANCHOR.MIDDLE)
        text(slide, ix + iw - 50, ry, 50, 32, [[(str(n), 17.5, C["text"], True, MONO)]],
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        ry += 33
    ry += 14
    shape(slide, ix, ry, iw, 1.5, C["border"])
    text(slide, ix, ry + 14, iw, 20, [[("THREE SCORER TYPES", 13.5, C["muted"], True, SANS, 150)]])
    ry += 44
    for title, detail in SCORERS:
        text(slide, ix, ry, iw, 50, [[(title, 17.5, C["white"], True)], [(detail, 15, C["muted"])]])
        ry += 54


def gates_panel(slide, x, y, w, h):
    panel(slide, x, y, w, h, "2", "Release quality gates", "PR gate = minutes, deterministic \u00b7 Nightly = LLM-judged")
    ix, iw = x + 28, w - 56
    header_h, row_h = 34, 47
    col_w = [iw - 140 - 120, 140, 120]
    gf = slide.shapes.add_table(1 + len(GATES), 3, px(ix), px(y + 104), px(iw), px(header_h + row_h * len(GATES)))
    tbl = gf.table
    tbl.first_row = tbl.horz_banding = False
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = px(cw)
    tbl.rows[0].height = px(header_h)
    for r in range(1, len(GATES) + 1):
        tbl.rows[r].height = px(row_h)

    def put(r, c, paragraphs, align):
        cell = tbl.cell(r, c)
        set_cell_borders(cell, bottom=C["border"] if r < len(GATES) else None)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(C["panel"])
        cell.margin_left = cell.margin_right = cell.margin_top = cell.margin_bottom = 0
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.text_frame.word_wrap = True
        fill_frame(cell.text_frame, paragraphs, align)

    for c, (head, al) in enumerate((("KPI", PP_ALIGN.LEFT), ("TARGET", PP_ALIGN.RIGHT), ("GATE", PP_ALIGN.RIGHT))):
        put(0, c, [[(head, 13.5, C["muted"], True, SANS, 150)]], al)
    for r, (kpi, target, gate) in enumerate(GATES, start=1):
        put(r, 0, [[(kpi, 18, C["text"])]], PP_ALIGN.LEFT)
        put(r, 1, [[(target, 18, C["white"], True, MONO)]], PP_ALIGN.RIGHT)
        put(r, 2, [[(gate, 16, C["cyan"] if gate == "PR" else C["pro"], True)]], PP_ALIGN.RIGHT)


def obs_panel(slide, x, y, w, h):
    panel(slide, x, y, w, h, "3", "Online observability", "Production telemetry on Google Cloud")
    ix, iw = x + 28, w - 56
    ry = y + 108
    for title, detail, color in OBSERVABILITY:
        shape(slide, ix, ry + 4, 5, 62, color)
        text(slide, ix + 20, ry, iw - 20, 70, [[(title, 17.5, C["white"], True)], [(detail, 15, C["muted"])]],
             line_spacing=1.05)
        ry += 82


def lifecycle(slide):
    y, h = 888, 140
    shape(slide, 72, y, 1776, h, "0D1119", C["border"], radius=14)
    text(slide, 96, y + 16, 600, 20, [[("LLMOPS LIFECYCLE \u00b7 CONTINUOUS QUALITY FLYWHEEL", 14, C["muted"], True, SANS, 200)]])
    n = len(LIFECYCLE)
    gap, x0, x1 = 34, 96, 1824
    w = (x1 - x0 - gap * (n - 1)) / n
    colors = ["94A3B8", "22D3EE", "C084FC", "38BDF8", "34D399", "FBBF24"]
    for i, (step, detail) in enumerate(LIFECYCLE):
        bx = x0 + i * (w + gap)
        shape(slide, bx, y + 48, w, 74, "141820", C["border"], radius=10)
        shape(slide, bx, y + 48, 5, 74, colors[i])
        text(slide, bx + 18, y + 54, w - 26, 64, [
            [(f"{i + 1}  ", 15, colors[i], True, MONO), (step, 18, C["white"], True)],
            [(detail, 14.5, C["muted"])],
        ], anchor=MSO_ANCHOR.MIDDLE)
        if i < n - 1:
            a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, px(bx + w + 8), px(y + 76), px(18), px(18))
            a.fill.solid()
            a.fill.fore_color.rgb = rgb(C["dim"])
            a.line.fill.background()


NOTES = """LLMOps, evaluation framework and observability (source: docs/EVALUATION_PLAN.md).

Offline: one versioned 100-ticket golden set taken from records in the RAG datastore, so every expected answer traces to a document ID. Composition: 45 search-ready in corpus, 25 urgent/frustrated, 15 vague single-turn, 10 multi-turn, 5 out of corpus. Scorers: deterministic metrics (routing F1, ROUGE-L, BLEU, search gating, latency, token accounting), an LLM judge (grounding recall, faithfulness) and per-case dynamic rubrics.

Gates: PR mode runs deterministic KPIs and blocks merges; nightly adds LLM-judged KPIs and failure analysis. Latency and error targets are proposed project targets (the TDD sets none).

Online: OpenTelemetry instrumentation of the FastAPI backend and httpx calls exports to Cloud Trace; Cloud Logging sinks to BigQuery Agent Analytics; sampled live traces are scored asynchronously; Model Armor / Cloud DLP redacts PII before it reaches agents or logs; the circuit breaker opens after 3 failures and retries after 30 s.

Status: framework and static rubric evaluator are built; the first baseline run has not happened yet, so no measured KPI values are shown.
"""


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    header(slide)
    y, h = 204, 660
    eval_panel(slide, 72, y, 540, h)
    gates_panel(slide, 636, y, 640, h)
    obs_panel(slide, 1300, y, 548, h)
    lifecycle(slide)
    slide.notes_slide.notes_text_frame.text = NOTES
    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
